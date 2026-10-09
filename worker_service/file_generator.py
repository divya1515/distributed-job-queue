import openpyxl
import os
from openpyxl.styles import Font,Alignment,PatternFill
from datetime import datetime
from pathlib import Path

FONT_NAME = "Arial"
RESULT_FILES_DIR = Path(os.getenv("JOB_FILES_DIR", "./job_files"))
RESULT_FILES_DIR.mkdir(parents=True, exist_ok=True)

def generate_job_result_file(job,job_items: list) -> str:
    wb = openpyxl.Workbook()
    summary=wb.active
    summary.title="Summary"
    summary["A1"] = "Job Summary"
    summary["A1"].font = Font(name=FONT_NAME, size=14, bold=True)

    summary_fields = [
        ("Job ID", str(job.job_id)),
        ("Status", job.status.value if hasattr(job.status, "value") else str(job.status)),
        ("Total Items", len(job_items)),
        ("Generated At", datetime.utcnow().strftime("%Y-%m-%d %H:%M:%S UTC")),
    ]
    for i, (label, value) in enumerate(summary_fields, start=3):
        summary[f"A{i}"] = label
        summary[f"A{i}"].font = Font(name=FONT_NAME, bold=True)
        summary[f"B{i}"] = value
        summary[f"B{i}"].font = Font(name=FONT_NAME)

    summary.column_dimensions["A"].width = 20
    summary.column_dimensions["B"].width = 30
    # ---- Details sheet ----
    details=wb.create_sheet("Details")
    headers = ["Original URL", "Short URL", "Status", "Attempts"]
    header_fill = PatternFill(start_color="DDDDDD", end_color="DDDDDD", fill_type="solid")
    for col, header in enumerate(headers, start=1):
        cell = details.cell(row=1, column=col, value=header)
        cell.font = Font(name=FONT_NAME, bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    for row, item in enumerate(job_items, start=2):
        details.cell(row=row, column=1, value=str(item.original_url)).font = Font(name=FONT_NAME)
        details.cell(row=row, column=2, value=str(item.short_url)).font = Font(name=FONT_NAME)
        details.cell(row=row, column=3, value=item.status).font = Font(name=FONT_NAME)
        details.cell(row=row, column=4, value=getattr(item, "attempts", "") or "").font = Font(name=FONT_NAME)

    for col_letter, width in zip("ABCD", [15, 15, 12, 40]):
        details.column_dimensions[col_letter].width = width

    file_path = RESULT_FILES_DIR / f"{job.job_id}.xlsx"
    wb.save(file_path)

    return str(file_path)